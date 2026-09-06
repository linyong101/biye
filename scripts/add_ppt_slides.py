#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
给「Vigil 答辩演示.pptx」追加两页（答辩现场演示 / 三大演示数据说明）。

用法（在本机终端，项目根目录 vigil/ 下）：
    pip install python-pptx      # 仅需首次
    python scripts/add_ppt_slides.py

说明：
- 默认生成新文件 docs/Vigil_答辩演示_增强.pptx，不破坏原文件；
- 若想直接覆盖原文件，把下方 OVERWRITE 改为 True 即可。
"""
import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError:
    print("缺少依赖：请先执行  pip install python-pptx")
    sys.exit(1)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "Vigil_答辩演示.pptx")
OVERWRITE = False

if not os.path.exists(SRC):
    print("未找到 PPT 文件：", SRC)
    sys.exit(1)

p = Presentation(SRC)

# 选择「标题 + 内容」版式（pptxgenjs 生成的标准模板里通常存在）
layout = None
for lo in p.slide_layouts:
    name = (lo.name or "").lower()
    if "title" in name and "content" in name:
        layout = lo
        break
if layout is None:
    layout = p.slide_layouts[1] if len(p.slide_layouts) > 1 else p.slide_layouts[0]


def add_slide(title, bullets):
    s = p.slides.add_slide(layout)
    s.shapes.title.text = title
    # 内容占位符（标题与内容版式中通常为 idx=1）
    body = s.placeholders.get(1)
    if body is None:
        # 兜底：直接加文本框
        from pptx.util import Inches
        box = s.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(9), Inches(5))
        body = box
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, b in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = b
        para.font.size = Pt(18)
        para.space_after = Pt(8)
    return s


add_slide(
    "答辩现场演示（30 秒跑通）",
    [
        "安装与初始化：npm install → npm run db:push（默认 SQLite，零依赖）",
        "一键填充演示数据：npm run seed（7 天错误/性能 + 会话回放剧本 + 三类告警规则）",
        "启动：npm run dev（看板 :5173）+ npm run demo（异常制造机 :5174）",
        "登录 admin / admin123；在 Demo 站点点「触发」即可在看板看到完整链路",
        "默认账号登录后请尽快修改密码",
    ],
)

add_slide(
    "一看即懂 · 三大演示数据已预置",
    [
        "概览页：错误趋势、TOP 问题、浏览器/版本分布、Web Vitals P75",
        "会话回放页：3 个剧本（购物车 500 报错 / 白屏 / 正常浏览），时间轴逐帧复现出错前操作",
        "告警配置页：new_issue / error_spike / perf_degrade 三类规则已预置，可现场讲解触发逻辑",
    ],
)

if OVERWRITE:
    dst = SRC
else:
    dst = os.path.join(ROOT, "docs", "Vigil_答辩演示_增强.pptx")

p.save(dst)
print("已生成增强版 PPT：", dst)
print("当前页数：", len(p.slides))
